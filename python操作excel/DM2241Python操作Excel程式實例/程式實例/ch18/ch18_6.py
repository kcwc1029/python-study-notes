# ch18_6.py
import openpyxl
from openpyxl.chart import ProjectedPieChart, Reference
from openpyxl.chart.series import DataPoint
from copy import deepcopy

wb = openpyxl.Workbook()                
ws = wb.active                          # 目前工作表
data = [
    ['產品','銷售業績'],
    ['化妝品', 85000],
    ['家電', 10000],
    ['日用品', 3000],
    ['文具', 2000],
]
for row in data:
    ws.append(row)
# 建立圓餅投影圖 --- pie
projected_pie = ProjectedPieChart()
projected_pie.type = "pie"              # 投影到 pie
projected_pie.splitType = "percent"     # 依百分比投影
# 設定資料來源
data = Reference(ws, min_col=2, min_row=1, max_row=5)
# 將資料加入圓餅投影圖物件
projected_pie.add_data(data, titles_from_data=True)
# 設定標籤資料
labels = Reference(ws,min_col=1,min_row=2,max_row=5)
projected_pie.set_categories(labels)
# 將圖表加入工作表
ws.add_chart(projected_pie, "D1")

# 建立圓餅投影圖 --- bar
projected_bar = deepcopy(projected_pie)
projected_bar.type = "bar"              # 投影到 bar
projected_bar.splitType = 'pos'         # 依位置投影
ws.add_chart(projected_bar, "D16")      # 將圖表加入工作表
wb.save('out18_6.xlsx')





