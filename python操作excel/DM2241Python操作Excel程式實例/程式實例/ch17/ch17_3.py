# ch17_3.py
import openpyxl
from openpyxl.chart import BubbleChart, Series
from openpyxl.chart import Reference

fn = "data17_3.xlsx"
wb = openpyxl.load_workbook(fn)
ws = wb.active

chart = BubbleChart()
chart.style = 48 
chart.title = ws['B2'].value

# 建立系列物件 s
# 建立 x 軸資料 xvalues
xvalues = Reference(ws,min_col=3,max_col=8,min_row=3)
# 建立 y 軸資料 yvalues
yvalues = Reference(ws,min_col=3,max_col=8,min_row=4)
# 建立 z 軸資料 size, 這是氣泡的大小
size = Reference(ws,min_col=3,max_col=8,min_row=5)
s = Series(values=yvalues,xvalues=xvalues,zvalues=size,
                title="2025年")
# 將系列物件 s 加入氣泡圖物件
chart.series.append(s)
# 將氣泡圖物件加入工作表, 放在 B7
ws.add_chart(chart,"B7")
wb.save("out17_3.xlsx")


